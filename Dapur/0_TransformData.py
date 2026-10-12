import os
import glob
import pandas as pd
import openpyxl

def clean_numeric(val):
    if pd.isna(val) or str(val).strip() in ['', 'nan', 'None']:
        return 0
    if isinstance(val, (int, float)):
        return val
    val_str = str(val).strip().replace('.', '').replace(',', '.')
    try:
        num = float(val_str)
        return int(num) if num.is_integer() else num
    except ValueError:
        return 0

def format_tanggal_faktur(val):
    if pd.isna(val) or str(val).strip() in ['', 'nan', 'None']:
        return ''
    val_str = str(val).strip()
    try:
        dt = pd.to_datetime(val_str, format='%d/%m/%Y')
        return dt.strftime('%Y-%m-%dT00:00:00')
    except (ValueError, TypeError):
        pass
    try:
        dt = pd.to_datetime(val_str)
        return dt.strftime('%Y-%m-%dT00:00:00')
    except (ValueError, TypeError):
        return val_str

def process_and_merge_excel():
    search_pattern = "data*.xlsx"
    file_list = glob.glob(search_pattern)
    output_filename = "data_export_123_temp.xlsx"
    file_list = [f for f in file_list if f != output_filename]
    if not file_list:
        print("--> Tidak ada file yang cocok dengan pola 'data*.xlsx' ditemukan.")
        return
    print(f"--> Ditemukan {len(file_list)} file untuk digabungkan.")
    text_columns = ['NPWP Pembeli / Identitas lainnya', 'Nomor Faktur Pajak', 'Referensi', 'Tanggal Faktur Pajak']
    numeric_columns = ['Harga Jual/Penggantian/DPP', 'DPP Nilai Lain/DPP', 'PPN', 'PPnBM']
    dfs = []
    for file in file_list:
        try:
            df = pd.read_excel(file, dtype=str)
            dfs.append(df)
            print(f"--> Berhasil membaca: {file}")
        except Exception as e:
            print(f"--> Gagal membaca {file}: {e}")
    if not dfs:
        print("--> Tidak ada data yang berhasil dibaca.")
        return
    combined_df = pd.concat(dfs, ignore_index=True)
    for col in ['NPWP Pembeli / Identitas lainnya', 'Nomor Faktur Pajak', 'Referensi']:
        if col in combined_df.columns:
            combined_df[col] = combined_df[col].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
            combined_df[col] = combined_df[col].replace(['nan', 'None'], '')
    if 'Tanggal Faktur Pajak' in combined_df.columns:
        combined_df['Tanggal Faktur Pajak'] = combined_df['Tanggal Faktur Pajak'].apply(format_tanggal_faktur)
    for col in numeric_columns:
        if col in combined_df.columns:
            combined_df[col] = combined_df[col].apply(clean_numeric)
    if 'IsShowClearName' not in combined_df.columns:
        combined_df['IsShowClearName'] = ''
    combined_df.insert(0, '', '')
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = list(combined_df.columns)
    ws.append(headers)
    numeric_col_indices = [i + 1 for i, col in enumerate(headers) if col in numeric_columns]
    text_col_indices = [i + 1 for i, col in enumerate(headers) if col in text_columns]
    for row in combined_df.itertuples(index=False):
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col_idx in numeric_col_indices:
            cell = row[col_idx - 1]
            cell.number_format = '#,##0'
        for col_idx in text_col_indices:
            cell = row[col_idx - 1]
            cell.number_format = '@'
    wb.save(output_filename)
    print(f"--> Selesai! Data berhasil digabungkan dan disimpan ke: {output_filename}")

if __name__ == "__main__":
    process_and_merge_excel()